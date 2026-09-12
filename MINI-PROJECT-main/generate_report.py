# ─────────────────────────────────────────────────────────────
#  AI BASED AUTOMATIC DOCUMENTATION GENERATOR — Report Builder
#  Run:  python generate_report.py
#  Output: AI_DOC_GENERATOR_REPORT.docx
# ─────────────────────────────────────────────────────────────
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page Setup (A4, academic margins) ───────────────────────
section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width  = Cm(21.0)
section.left_margin   = Cm(3.0)
section.right_margin  = Cm(2.0)
section.top_margin    = Cm(2.5)
section.bottom_margin = Cm(2.5)

# ── Helper Functions ─────────────────────────────────────────
TNR = "Times New Roman"

def _set_run(run, size=12, bold=False, italic=False, underline=False, color=None):
    run.font.name = TNR
    run.font.size = Pt(size)
    run.bold      = bold
    run.italic    = italic
    run.underline = underline
    if color:
        run.font.color.rgb = RGBColor(*color)

def _set_para(para, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=8, line_spacing=1.5):
    pf = para.paragraph_format
    pf.alignment     = align
    pf.space_before  = Pt(space_before)
    pf.space_after   = Pt(space_after)
    pf.line_spacing  = Pt(line_spacing * 12)

def add_page_break():
    doc.add_page_break()

def heading1(text):
    """Chapter title — centered, bold, 16pt, uppercase"""
    p = doc.add_paragraph()
    _set_para(p, WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=12)
    r = p.add_run(text.upper())
    _set_run(r, size=16, bold=True)
    return p

def heading2(text):
    """Section — left, bold, underline, 14pt"""
    p = doc.add_paragraph()
    _set_para(p, WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=6)
    r = p.add_run(text)
    _set_run(r, size=14, bold=True, underline=True)
    return p

def heading3(text):
    """Sub-section — left, bold, 12pt"""
    p = doc.add_paragraph()
    _set_para(p, WD_ALIGN_PARAGRAPH.LEFT, space_before=8, space_after=4)
    r = p.add_run(text)
    _set_run(r, size=12, bold=True)
    return p

def body(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    _set_para(p, align, space_after=8)
    r = p.add_run(text)
    _set_run(r, size=12)
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    _set_para(p, WD_ALIGN_PARAGRAPH.LEFT, space_after=4)
    r = p.add_run(text)
    _set_run(r, size=12)
    return p

def numbered(text):
    p = doc.add_paragraph(style="List Number")
    _set_para(p, WD_ALIGN_PARAGRAPH.LEFT, space_after=4)
    r = p.add_run(text)
    _set_run(r, size=12)
    return p

def center_text(text, size=12, bold=False):
    p = doc.add_paragraph()
    _set_para(p, WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    r = p.add_run(text)
    _set_run(r, size=size, bold=bold)
    return p

def add_table(headers, rows, col_widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    hrow = t.rows[0]
    for i, h in enumerate(headers):
        cell = hrow.cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        _set_run(run, size=11, bold=True)
        cell.paragraphs[0].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "D9D9D9")
        shd.set(qn("w:val"), "clear")
        cell._tc.get_or_add_tcPr().append(shd)
    for ri, row_data in enumerate(rows):
        row = t.rows[ri + 1]
        for ci, val in enumerate(row_data):
            cell = row.cells[ci]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            _set_run(run, size=11)
            cell.paragraphs[0].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if col_widths:
        for ri2, row in enumerate(t.rows):
            for ci2, cell in enumerate(row.cells):
                if ci2 < len(col_widths):
                    cell.width = Cm(col_widths[ci2])
    doc.add_paragraph()
    return t

def figure_placeholder(fig_num, caption):
    """Inserts a bordered box as diagram placeholder"""
    t = doc.add_table(rows=3, cols=1)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = ""
    t.rows[0].cells[0].paragraphs[0].add_run(f"Figure {fig_num}").bold = True
    c = t.rows[1].cells[0]
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"[ {caption} ]")
    _set_run(r, size=11, italic=True)
    t.rows[1].height = Cm(5)
    t.rows[2].cells[0].text = ""
    cp = t.rows[2].cells[0].paragraphs[0]
    cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cp.add_run(f"Figure {fig_num}: {caption}")
    _set_run(cr, size=11, italic=True, bold=True)
    doc.add_paragraph()

# ════════════════════════════════════════════════════════════
#  COVER PAGE
# ════════════════════════════════════════════════════════════
center_text("AI BASED AUTOMATIC DOCUMENTATION GENERATOR", size=16, bold=True)
center_text("MINI PROJECT REPORT", size=14, bold=True)
doc.add_paragraph()
center_text("Submitted by", size=12)
doc.add_paragraph()

students = [
    ("BALAMANIGANDAN S", "142223104020"),
    ("BALAMURUGAN M",    "142223104022"),
    ("DEVAPRASANTH G",   "142223104025"),
]
t_cover = doc.add_table(rows=len(students), cols=2)
t_cover.style = "Table Grid"
for i, (name, roll) in enumerate(students):
    t_cover.rows[i].cells[0].text = name
    t_cover.rows[i].cells[1].text = roll
    for cell in t_cover.rows[i].cells:
        cell.paragraphs[0].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            _set_run(run, size=12, bold=True)
doc.add_paragraph()
center_text("in partial fulfilment for the award of the degree", size=12)
center_text("of", size=12)
center_text("BACHELOR OF ENGINEERING", size=13, bold=True)
center_text("in", size=12)
center_text("COMPUTER SCIENCE AND ENGINEERING", size=13, bold=True)
doc.add_paragraph()
center_text("SRM VALLIAMMAI ENGINEERING COLLEGE", size=13, bold=True)
center_text("(AN AUTONOMOUS INSTITUTION)", size=12)
center_text("SRM NAGAR, KATTANKULATHUR, CHENGALPATTU", size=12)
center_text("ANNA UNIVERSITY: CHENNAI 600 025", size=12)
doc.add_paragraph()
center_text("JUNE 2026", size=13, bold=True)
add_page_break()

# ════════════════════════════════════════════════════════════
#  BONAFIDE CERTIFICATE
# ════════════════════════════════════════════════════════════
heading1("BONAFIDE CERTIFICATE")
body(
    'Certified that this project report "AI BASED AUTOMATIC DOCUMENTATION GENERATOR" '
    'is the Bonafide work of "BALAMANIGANDAN S (142223104020), BALAMURUGAN M (142223104022), '
    'DEVAPRASANTH G (142223104025)" who carried out the project work under my supervision '
    'during the academic year 2025-26.'
)
doc.add_paragraph()

sig_table = doc.add_table(rows=4, cols=2)
sig_table.style = "Table Grid"
sig_data = [
    ("SIGNATURE", "SIGNATURE"),
    ("[Supervisor Name]", "Dr. B. VANATHI, M.E., Ph.D.,"),
    ("[Designation], Department of CSE", "Professor & Head, Department of CSE"),
    ("SRM Valliammai Engineering College, Kattankulathur – 603 203.",
     "SRM Valliammai Engineering College, Kattankulathur – 603 203."),
]
for ri, (l, r_) in enumerate(sig_data):
    sig_table.rows[ri].cells[0].text = l
    sig_table.rows[ri].cells[1].text = r_
    for cell in sig_table.rows[ri].cells:
        for run in cell.paragraphs[0].runs:
            _set_run(run, size=12, bold=(ri == 0))
doc.add_paragraph()
body("Submitted for the University examination held on __________ at SRM Valliammai Engineering College, Kattankulathur.")
doc.add_paragraph()
p_ex = doc.add_paragraph()
_set_para(p_ex, WD_ALIGN_PARAGRAPH.LEFT)
p_ex.add_run("INTERNAL EXAMINER" + " " * 30 + "EXTERNAL EXAMINER").bold = True
add_page_break()

# ════════════════════════════════════════════════════════════
#  ACKNOWLEDGEMENT
# ════════════════════════════════════════════════════════════
heading1("ACKNOWLEDGEMENT")
body("We sincerely express our gratitude in depth to our esteemed Founder Chairman & Chancellor Dr.T.R.Paarivendhar, Thiru.Ravi Pachamoothoo Chairman, Mrs. Padmapriya Ravi Vice Chairman, Ms.R.Harini Correspondent, SRM VALLIAMMAI ENGINEERING COLLEGE for the patronage on our welfare rooted in the academic year.")
body("We express our sincere gratitude to our respected Dr.Chidhambararajan, M.E., Ph.D., for his constant encouragement, which has been our motivation to strive towards excellence.")
body("We thank our Principal Dr.M.Murugan, M.E., Ph.D., for his constant encouragement, which has been our motivation to strive towards excellence.")
body("We also thank our Vice Principal Dr.S.Visalakshi Selvaraj, M.E., Ph.D., for her constant motivation to strive towards excellence.")
body("We extend our hand of thanks to our Head of the Department, Dr.B.Vanathi, M.E., Ph.D., Professor for her unstinted support.")
body("We are thankful to our project coordinator, [Supervisor Name], [Designation], for providing us with the essential facilities and guiding us at every step of this project with patience and expertise.")
body("We also like to thank all Teaching and non-teaching staff members of our department for their support during the course of the project. We finally thank our friends and family for their unwavering support and encouragement throughout this journey.")
add_page_break()

# ════════════════════════════════════════════════════════════
#  ABSTRACT — ENGLISH
# ════════════════════════════════════════════════════════════
heading1("ABSTRACT")
body("Software documentation forms a critical backbone of every software engineering project, serving as the primary reference for architecture, functionality, and system maintenance. Despite its immense importance, generating comprehensive academic documentation remains a tedious, time-intensive, and error-prone process for students and developers alike. Compiling a standard final-year academic project report requires meticulous formatting, careful reverse-engineering of implemented codebases, manual creation of architectural diagrams, and strict adherence to institutional formatting guidelines. This process typically consumes several weeks of effort that could otherwise be directed toward improving the software itself.")
body("To address these challenges, this project presents the AI Based Automatic Documentation Generator — an innovative, intelligent web application designed to fully automate the academic documentation pipeline. The platform directly interfaces with GitHub repositories via the Octokit REST API, ingests source code files, and performs deep structural analysis using Abstract Syntax Tree (AST) parsing with @babel/parser. Key code artifacts including classes, functions, API endpoints, database schemas (Prisma), environment variable configurations, and file roles are automatically extracted and compiled into a rich project context.")
body("Built on the Next.js 16 App Router framework and secured with Clerk authentication, the system employs an intelligent multi-model AI routing engine (ModelRouter) that dispatches twelve distinct documentation task types to the most suitable AI provider. Google Gemini 1.5 Pro serves as the free-tier primary model, with OpenAI GPT-4o as an optional secondary provider and a local Ollama instance (gemma2) as a final offline fallback. Each AI task is configured with a specifically tuned temperature value — lower values (0.2) for structured code analysis and diagram generation, and higher values (0.7) for creative academic prose writing — ensuring both technical precision and linguistic quality.")
body("The system automatically generates four Mermaid.js diagrams — System Architecture, Component, API Sequence, and Entity-Relationship — by encoding them into mermaid.ink SVG URLs that are embedded directly into the document. A custom marked.js renderer transforms the AI-generated Markdown into semantically correct HTML, which is then rendered into a publication-ready, A4-formatted PDF using Puppeteer with academic margins of 30mm left, 20mm right, and 25mm top and bottom. The final output includes a structured eight-chapter academic report complete with cover page, bonafide certificate, abstract, table of contents, and references.")
body("Testing and validation across multiple real-world GitHub repositories confirmed that the system successfully reduces academic report generation time from several weeks to under two minutes, while maintaining high linguistic quality, technical accuracy, and strict formatting compliance. The result is a powerful, cost-effective, and privacy-conscious tool that empowers students to focus on software development rather than documentation overhead.")
doc.add_paragraph()
kw = doc.add_paragraph()
_set_para(kw, WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)
kr = kw.add_run("KEYWORDS: ")
_set_run(kr, size=12, bold=True)
kr2 = kw.add_run("Artificial Intelligence, Automatic Documentation, GitHub API, Gemini 1.5 Pro, Next.js 16, Puppeteer, AST Parsing, Mermaid.js, Multi-Model Routing, Academic Report Generator, Clerk Authentication, Ollama, PDF Generation.")
_set_run(kr2, size=12)
add_page_break()

# ════════════════════════════════════════════════════════════
#  ABSTRACT — TAMIL
# ════════════════════════════════════════════════════════════
heading1("சுருக்கம்")
body("மென்பொருள் ஆவண தயாரிப்பு என்பது மென்பொருள் பொறியியலில் மிக முக்கியமான ஒரு பணியாகும். ஆனால் இது மாணவர்களுக்கும் உருவாக்குநர்களுக்கும் மிகவும் நேரத்தை வீணடிக்கும் மற்றும் சிரமகரமான ஒரு செயலாக உள்ளது. பாரம்பரிய ஆவண தயாரிப்பு முறைகள் கைமுறை முயற்சிகளை அதிகமாக தேவைப்படுகின்றன மற்றும் தவறுகளுக்கு வழிவகுக்கின்றன. ஒரு இறுதி ஆண்டு கல்வி திட்ட அறிக்கையை தயாரிக்க பல வாரங்கள் தேவைப்படும், இந்த நேரம் மென்பொருள் தரத்தை மேம்படுத்துவதில் செலவிடப்பட வேண்டும்.")
body("இந்த சிக்கல்களை தீர்க்க, இந்த திட்டம் 'AI அடிப்படையிலான தானியங்கி ஆவண உருவாக்கி' என்ற ஒரு புதுமையான இணைய பயன்பாட்டை வழங்குகிறது. இந்த அமைப்பு GitHub களஞ்சியங்களை Octokit REST API மூலம் நேரடியாக இணைக்கிறது மற்றும் @babel/parser கருவியைப் பயன்படுத்தி Abstract Syntax Tree (AST) பகுப்பாய்வு மூலம் குறியீட்டை ஆழமாக ஆய்வு செய்கிறது. வகுப்புகள், செயல்பாடுகள், API இறுதிப்புள்ளிகள், தரவுத்தள திட்டங்கள் மற்றும் கோப்பு பாத்திரங்கள் ஆகியவை தானாகவே கண்டறியப்படுகின்றன.")
body("Next.js 16 கட்டமைப்பில் கட்டமைக்கப்பட்ட இந்த அமைப்பு, Clerk அங்கீகார அமைப்பு மூலம் பாதுகாக்கப்படுகிறது. ModelRouter என்ற சிறப்பு AI வழிசெலுத்தி, Google Gemini 1.5 Pro ஐ முதன்மை மாதிரியாகவும், OpenAI GPT-4o ஐ இரண்டாம் நிலை மாதிரியாகவும், உள்ளூர் Ollama (gemma2) ஐ இறுதி மாற்று வழியாகவும் பயன்படுத்தி, பன்னிரண்டு வகையான ஆவண பணிகளை சரியான மாதிரிக்கு திருப்பி விடுகிறது.")
body("Mermaid.js தொழில்நுட்பத்தைப் பயன்படுத்தி கட்டமைப்பு வரைபடம், கூறு வரைபடம், வரிசை வரைபடம் மற்றும் உறவு வரைபடம் ஆகிய நான்கு வரைபடங்கள் தானியங்கியாக உருவாக்கப்படுகின்றன. marked.js மற்றும் Puppeteer மூலம் A4 வடிவமைப்பில் தொழில்முறை PDF ஆவணங்கள் உருவாக்கப்படுகின்றன. இந்த அமைப்பு மாணவர்கள் பல வாரங்கள் செலவிடும் ஆவண தயாரிப்பு நேரத்தை இரண்டு நிமிடங்களுக்கும் குறைவாக குறைக்கிறது.")
doc.add_paragraph()
tkw = doc.add_paragraph()
_set_para(tkw, WD_ALIGN_PARAGRAPH.JUSTIFY)
tkr = tkw.add_run("முக்கிய வார்த்தைகள்: ")
_set_run(tkr, size=12, bold=True)
tkr2 = tkw.add_run("செயற்கை நுண்ணறிவு, தானியங்கி ஆவண உருவாக்கி, GitHub API, Gemini 1.5 Pro, Next.js 16, Puppeteer, AST பகுப்பாய்வு, Mermaid.js, பல மாதிரி வழிசெலுத்தல், கல்வி அறிக்கை உருவாக்கம்.")
_set_run(tkr2, size=12)
add_page_break()

# ════════════════════════════════════════════════════════════
#  TABLE OF CONTENTS
# ════════════════════════════════════════════════════════════
heading1("TABLE OF CONTENTS")
toc_data = [
    ("S.NO", "TITLE", "PAGE NO."),
    ("—", "BONAFIDE CERTIFICATE", "i"),
    ("—", "ACKNOWLEDGEMENT", "ii"),
    ("—", "ABSTRACT (ENGLISH)", "iii"),
    ("—", "ABSTRACT (TAMIL)", "iv"),
    ("—", "TABLE OF CONTENTS", "v"),
    ("—", "LIST OF TABLES", "vi"),
    ("—", "LIST OF FIGURES", "vii"),
    ("1", "INTRODUCTION", "1"),
    ("1.1", "Project Introduction", "1"),
    ("1.2", "Problem Statement", "3"),
    ("1.3", "Overview of the Project", "4"),
    ("1.4", "Objectives", "5"),
    ("1.5", "Scope of the Project", "6"),
    ("1.6", "Motivation", "7"),
    ("1.7", "Organization of the Report", "8"),
    ("2", "LITERATURE SURVEY", "9"),
    ("3", "METHODOLOGY", "14"),
    ("3.1", "Methodology Overview", "14"),
    ("3.2", "Stage-wise Methodology", "14"),
    ("3.3", "System Requirements", "18"),
    ("4", "SYSTEM DESIGN", "20"),
    ("4.1", "System Architecture", "20"),
    ("4.2", "Flow Diagram", "21"),
    ("4.3", "Use Case Diagram", "22"),
    ("4.4", "Sequence Diagram", "23"),
    ("4.5", "Activity Diagram", "24"),
    ("4.6", "Data Flow Diagram", "25"),
    ("5", "IMPLEMENTATION AND RESULTS", "27"),
    ("5.1", "Implementation Overview", "27"),
    ("5.2", "Module-wise Implementation", "28"),
    ("5.3", "Output and Results", "38"),
    ("6", "TESTING AND VALIDATION", "40"),
    ("6.1", "Testing Methodology", "40"),
    ("6.2", "Types of Testing", "40"),
    ("6.3", "Test Cases", "43"),
    ("6.4", "Testing Results", "45"),
    ("6.5", "Validation", "46"),
    ("6.6", "Validation Results", "48"),
    ("7", "CONCLUSION AND FUTURE SCOPE", "50"),
    ("7.1", "Conclusion", "50"),
    ("7.2", "Future Scope", "52"),
    ("—", "REFERENCES", "54"),
    ("—", "SDG CERTIFICATE", "56"),
]
t_toc = doc.add_table(rows=len(toc_data), cols=3)
t_toc.style = "Table Grid"
for ri, row in enumerate(toc_data):
    for ci, val in enumerate(row):
        cell = t_toc.rows[ri].cells[ci]
        cell.text = ""
        r = cell.paragraphs[0].add_run(val)
        is_hdr = (ri == 0)
        _set_run(r, size=11, bold=is_hdr)
        align = WD_ALIGN_PARAGRAPH.CENTER if ci in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
        cell.paragraphs[0].paragraph_format.alignment = align
        if is_hdr:
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "D9D9D9")
            shd.set(qn("w:val"), "clear")
            cell._tc.get_or_add_tcPr().append(shd)
for row in t_toc.rows:
    row.cells[0].width = Cm(2)
    row.cells[1].width = Cm(12)
    row.cells[2].width = Cm(2.5)
add_page_break()

print("Part 1 done — front matter complete.")
